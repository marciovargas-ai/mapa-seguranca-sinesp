"""
Baixa a base pública do Sinesp VDE (Ministério da Justiça e Segurança Pública),
resume por UF, município, evento e mês e grava um JSON pequeno por ano e UF:

  dados/sinesp_<ANO>_<UF>.json
  { v, ano, uf, importadoEm, fonte,
    estado:     { evento: [12 meses] },
    municipios: { nome_normalizado: { n: nome, e: { evento: [12 meses] } } } }
  dados/manifesto.json  { atualizadoEm, anos: { "<ANO>": { publicadoAte, ufs } } }

O Apps Script do Mapa da Segurança lê esses arquivos pelo endereço "raw" do
GitHub. Os arquivos de cada ano têm ~26 MB e ~750 mil linhas, grandes demais
para o Google Sheets e para o tempo de execução do Apps Script.

Uso: python scripts/atualizar_sinesp.py [--anos 2024 2025 2026] [--arquivo ano=caminho.xlsx]
"""
import argparse
import datetime as dt
import io
import json
import os
import re
import sys
import time
import unicodedata
import urllib.request
import zipfile

from openpyxl import load_workbook

URL_VDE = ('https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/'
           'download/dnsp-base-de-dados/bancovde-{ano}.xlsx/@@download/file')
PASTA_SAIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'dados')

UFS = {
    'AC': 'Acre', 'AL': 'Alagoas', 'AP': 'Amapá', 'AM': 'Amazonas', 'BA': 'Bahia', 'CE': 'Ceará',
    'DF': 'Distrito Federal', 'ES': 'Espírito Santo', 'GO': 'Goiás', 'MA': 'Maranhão', 'MT': 'Mato Grosso',
    'MS': 'Mato Grosso do Sul', 'MG': 'Minas Gerais', 'PA': 'Pará', 'PB': 'Paraíba', 'PR': 'Paraná',
    'PE': 'Pernambuco', 'PI': 'Piauí', 'RJ': 'Rio de Janeiro', 'RN': 'Rio Grande do Norte',
    'RS': 'Rio Grande do Sul', 'RO': 'Rondônia', 'RR': 'Roraima', 'SC': 'Santa Catarina', 'SP': 'São Paulo',
    'SE': 'Sergipe', 'TO': 'Tocantins'
}
MESES = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro',
         'outubro', 'novembro', 'dezembro']


def normalizar(texto):
    """Mesma regra de normalizarTexto_ no Apps Script (4_Utils_Gatilhos.gs)."""
    s = unicodedata.normalize('NFD', str(texto or '').lower())
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    s = re.sub(r'[^a-z0-9\s]', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


UF_POR_NOME = {}
for sigla, nome in UFS.items():
    UF_POR_NOME[normalizar(nome)] = sigla
    UF_POR_NOME[sigla.lower()] = sigla


def baixar(ano, tentativas=4):
    """Baixa o .xlsx do ano. O gov.br às vezes corta o arquivo com HTTP 200: valida o zip."""
    url = URL_VDE.format(ano=ano)
    for i in range(tentativas):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (mapa-da-seguranca)'})
            with urllib.request.urlopen(req, timeout=300) as resp:
                dados = resp.read()
            with zipfile.ZipFile(io.BytesIO(dados)) as z:
                if z.testzip() is None:
                    return dados
            print(f'{ano}: arquivo corrompido, nova tentativa')
        except urllib.error.HTTPError as e:
            if e.code == 404:
                print(f'{ano}: ainda não publicado')
                return None
            print(f'{ano}: HTTP {e.code}, nova tentativa')
        except Exception as e:  # noqa: BLE001 - rede instável: tenta de novo
            print(f'{ano}: {e}, nova tentativa')
        time.sleep(10 * (i + 1))
    return None


def indice_mes(valor, ano):
    if isinstance(valor, (dt.datetime, dt.date)):
        return valor.month - 1 if valor.year == ano else -1
    s = str(valor or '').strip()
    m = re.match(r'^(\d{4})-(\d{1,2})', s)
    if m:
        return int(m.group(2)) - 1 if int(m.group(1)) == ano else -1
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        return int(m.group(2)) - 1 if int(m.group(3)) == ano else -1
    m = re.match(r'^(\d{1,2})/(\d{4})', s)
    if m:
        return int(m.group(1)) - 1 if int(m.group(2)) == ano else -1
    t = normalizar(s)
    if t.isdigit() and 1 <= int(t) <= 12:
        return int(t) - 1
    for i, nome in enumerate(MESES):
        if t.startswith(nome) or (len(t) >= 3 and nome.startswith(t[:3])):
            return i
    return -1


def achar_colunas(cabecalho):
    nomes = [normalizar(v).replace(' ', '_') for v in cabecalho]

    def achar(*alvos):
        for i, n in enumerate(nomes):
            if any(n == a or n.startswith(a) for a in alvos):
                return i
        return -1

    c = {
        'uf': achar('uf', 'sigla_uf', 'estado'),
        'mun': achar('municipio', 'nome_municipio', 'cidade'),
        'evento': achar('evento', 'tipo_crime', 'indicador', 'natureza'),
        'data': achar('data_referencia', 'data', 'mes_ano', 'referencia'),
        'total': next((i for i, n in enumerate(nomes) if n in ('total', 'quantidade', 'ocorrencias', 'qtd')), -1),
        'vitimas': next((i for i, n in enumerate(nomes) if n in ('total_vitima', 'total_vitimas', 'vitimas')), -1),
        'abrangencia': achar('abrangencia'),
    }
    ok = c['uf'] >= 0 and c['evento'] >= 0 and c['data'] >= 0 and (c['total'] >= 0 or c['vitimas'] >= 0)
    return c if ok else None


def doze():
    return [0] * 12


def processar(conteudo, ano):
    """Devolve {uf: {estado, municipios, publicadoAte}} e o número de linhas lidas."""
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    est_estadual, est_municipal, municipios = {}, {}, {}
    ultimo_mes = {}
    lidas = 0
    for aba in wb.worksheets:
        col = None
        for linha in aba.iter_rows(values_only=True):
            if col is None:
                col = achar_colunas(linha)
                continue
            uf = UF_POR_NOME.get(normalizar(linha[col['uf']]))
            evento = str(linha[col['evento']] or '').strip()
            if not uf or not evento:
                continue
            mes = indice_mes(linha[col['data']], ano)
            if mes < 0:
                continue
            qtd = linha[col['total']] if col['total'] >= 0 else 0
            if not qtd and col['vitimas'] >= 0:
                qtd = linha[col['vitimas']]
            try:
                qtd = float(qtd or 0)
            except (TypeError, ValueError):
                qtd = 0
            if not qtd:
                continue
            qtd = int(qtd) if qtd == int(qtd) else qtd
            lidas += 1
            ultimo_mes[uf] = max(ultimo_mes.get(uf, -1), mes)

            nome_mun = str(linha[col['mun']] or '').strip() if col['mun'] >= 0 else ''
            abr = str(linha[col['abrangencia']] or '') if col['abrangencia'] >= 0 else ''
            estadual = bool(re.search(r'estad', abr, re.I)) or not nome_mun or re.search(r'n[aã]o informado', nome_mun, re.I)
            if estadual:
                est_estadual.setdefault(uf, {}).setdefault(evento, doze())[mes] += qtd
            else:
                est_municipal.setdefault(uf, {}).setdefault(evento, doze())[mes] += qtd
                chave = normalizar(nome_mun)
                m = municipios.setdefault(uf, {}).setdefault(chave, {'n': nome_mun.title(), 'e': {}})
                m['e'].setdefault(evento, doze())[mes] += qtd
        if col is None:
            print(f'{ano}: aba "{aba.title}" sem as colunas esperadas')
    wb.close()

    # Estado: a linha estadual quando existe; senão a soma dos municípios (evita contar duas vezes)
    saida = {}
    for uf in set(est_estadual) | set(est_municipal):
        eventos = {}
        for ev in set(est_estadual.get(uf, {})) | set(est_municipal.get(uf, {})):
            s = est_estadual.get(uf, {}).get(ev)
            eventos[ev] = s if s and any(s) else est_municipal[uf][ev]
        saida[uf] = {'estado': eventos, 'municipios': municipios.get(uf, {}), 'publicadoAte': ultimo_mes.get(uf, -1)}
    return saida, lidas


def gravar(caminho, obj):
    with open(caminho, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, separators=(',', ':'))


def main():
    ap = argparse.ArgumentParser()
    hoje = dt.date.today()
    ap.add_argument('--anos', nargs='*', type=int, default=[hoje.year - 2, hoje.year - 1, hoje.year])
    ap.add_argument('--arquivo', action='append', default=[], help='ano=caminho.xlsx (usa um arquivo local)')
    args = ap.parse_args()
    locais = dict(a.split('=', 1) for a in args.arquivo)

    os.makedirs(PASTA_SAIDA, exist_ok=True)
    caminho_manifesto = os.path.join(PASTA_SAIDA, 'manifesto.json')
    manifesto = {'v': 1, 'anos': {}}
    if os.path.exists(caminho_manifesto):
        with open(caminho_manifesto, encoding='utf-8') as f:
            manifesto = json.load(f)

    agora = int(time.time() * 1000)
    mudou = False
    for ano in args.anos:
        if str(ano) in locais:
            with open(locais[str(ano)], 'rb') as f:
                conteudo = f.read()
        else:
            conteudo = baixar(ano)
        if not conteudo:
            continue  # mantém o que já foi publicado
        resumo, lidas = processar(conteudo, ano)
        if not lidas:
            print(f'{ano}: nenhuma linha reconhecida; arquivos anteriores mantidos')
            continue
        for uf, a in resumo.items():
            gravar(os.path.join(PASTA_SAIDA, f'sinesp_{ano}_{uf}.json'), {
                'v': 3, 'ano': ano, 'uf': uf, 'importadoEm': agora, 'fonte': f'Sinesp VDE {ano} (gov.br)',
                'estado': a['estado'], 'municipios': a['municipios']
            })
        manifesto['anos'][str(ano)] = {
            'publicadoAte': max(a['publicadoAte'] for a in resumo.values()) + 1,
            'ufs': sorted(resumo), 'linhas': lidas, 'importadoEm': agora
        }
        mudou = True
        print(f'{ano}: {lidas} linhas, {len(resumo)} UFs')

    if mudou:
        manifesto['atualizadoEm'] = agora
        gravar(caminho_manifesto, manifesto)
    elif not manifesto['anos']:
        sys.exit('Nenhum ano foi processado.')


if __name__ == '__main__':
    main()
